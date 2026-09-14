import re
import time
import openpyxl
import cloudscraper
from bs4 import BeautifulSoup
from Base import get_excel_path


def fetch_page(url: str) -> str | None:
    """Fetch a Chittorgarh IPO page and return HTML text."""
    scraper = cloudscraper.create_scraper(
        browser={"browser": "chrome", "platform": "windows", "desktop": True}
    )
    try:
        response = scraper.get(url, timeout=20)
        response.raise_for_status()
        return response.text
    except Exception as e:
        print(f"  Error fetching page: {e}")
        return None


def extract_listing_price(html: str) -> float | None:
    """Extract listing open price from 'Listing Day Trading Information' table."""
    soup = BeautifulSoup(html, "html.parser")
    tables = soup.find_all("table")
    for table in tables:
        rows = table.find_all("tr")
        for row in rows:
            cells = row.find_all(["td", "th"])
            if len(cells) >= 2:
                label = cells[0].get_text(strip=True).lower()
                if label == "open":
                    for cell in cells[1:]:
                        text = cell.get_text(strip=True)
                        nums = re.findall(r"[\d,]+\.?\d*", text)
                        if nums:
                            try:
                                return float(nums[0].replace(",", ""))
                            except ValueError:
                                continue
    return None


def extract_issue_price(html: str) -> float | None:
    """Extract issue price from Chittorgarh IPO page HTML.

    Strategy:
    1. Parse with BeautifulSoup, get clean text (strips HTML tags)
    2. Find 'Issue Price' followed by rupee symbol and number + 'per share'
       (this is the real issue price in the IPO Details table)
    3. Skip the navigation link 'Issue price v/s Market price'
    4. Fallback: look for 'set final issue price at Rs.XX'
    """
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    text = soup.get_text()

    # Pattern 1 (best): "Issue Price₹XXX per share" in IPO Details table
    # We search the clean text, not raw HTML, to avoid tag interference
    patterns = [
        r"Issue Price[^₹\d]*₹\s*([\d,]+\.?\d*)\s*per share",
        r"Issue Price[^₹\d]*₹\s*([\d,]+\.?\d*)\s*(?:per equity|per share|per scrip)",
        r"set final issue price at[^₹\d]*₹\s*([\d,]+\.?\d*)",
        r"Issue Price[^₹\d]*Rs\.?\s*([\d,]+\.?\d*)\s*per share",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if m:
            try:
                val = float(m.group(1).replace(",", ""))
                if val > 1:  # Sanity check: skip placeholders
                    return val
            except ValueError:
                continue
    return None


def update_listing_results():
    path = get_excel_path()
    wb = openpyxl.load_workbook(path)
    sheets = ["IPOSME", "IPOMB"]
    total_updated = 0
    total_skipped = 0

    for sheet_name in sheets:
        ws = wb[sheet_name]
        print(f"\nProcessing sheet: {sheet_name}")
        updated = 0

        for row in range(2, ws.max_row + 1):
            company = ws.cell(row, 2).value
            url = ws.cell(row, 1).value
            d_val = ws.cell(row, 4).value

            # Skip empty rows or rows already filled (1 or 0)
            if not company or not url:
                continue
            if d_val in (0, 1):
                continue

            print(f"Row {row}: {company}")

            # Try reading issue price from Excel first (col 45)
            issue_price = None
            raw_issue = ws.cell(row, 45).value
            try:
                ip = float(raw_issue)
                if ip > 1:  # Real value, not a placeholder
                    issue_price = ip
            except (TypeError, ValueError):
                pass

            # Fetch Chittorgarh page once to get both prices
            html = fetch_page(url)
            time.sleep(1)

            if html is None:
                print(f"  Skipped: could not fetch page")
                total_skipped += 1
                continue

            # Extract listing price from page
            listing_price = extract_listing_price(html)

            # Fallback: get issue price from page if not in Excel
            if issue_price is None:
                issue_price = extract_issue_price(html)

            if listing_price is None or issue_price is None:
                print(f"  Skipped: listing={listing_price}, issue={issue_price}")
                total_skipped += 1
                continue

            result = 1 if listing_price >= issue_price else 0
            ws.cell(row, 4, result)
            updated += 1
            print(f"  listing={listing_price}, issue={issue_price} => {result}")

        wb.save(path)
        total_updated += updated
        print(f"Sheet {sheet_name}: {updated} rows updated")

    print(f"\nDone. Total rows updated: {total_updated}, Skipped: {total_skipped}")


if __name__ == "__main__":
    update_listing_results()
