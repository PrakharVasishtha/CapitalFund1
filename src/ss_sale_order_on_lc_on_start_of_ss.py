"""
ss_sale_order_on_lc_on_start_of_ss.py
======================================
Places a lower circuit (LC) sell order on Zerodha Kite at the start of the special pre-open session
for all newly allotted shares today (special_session_status == 5 in allotted_holdings.xlsx) across all users.

Upon successful placement of the sell order:
  - Updates special_session_status in allotted_holdings.xlsx to 1 (order placed).
  - Saves the updated Excel workbook.

Usage:
  python src/ss_sale_order_on_lc_on_start_of_ss.py
"""
import sys
import os

sys.path.append(os.path.dirname(os.path.abspath(__file__)))

import openpyxl
import Base
from config import SpecialSessionStatus
from allotment_update import get_allotted_holdings_path
from special_sesion_zerodha_sell import zerodha_sell_lc
from logger_setup import get_logger

logger = get_logger(__name__)

CREDENTIALS_FILE = "credentials.json"


def place_lc_sell_orders_for_allotted_today():
    """
    Scans allotted_holdings.xlsx for rows with special_session_status == 5 (allotted today)
    across all user sheets and triggers lower circuit sell orders on Zerodha Kite via zerodha_sell_lc().
    """
    logger.info("-----------place_lc_sell_orders_for_allotted_today----------")
    users = Base.load_credentials(CREDENTIALS_FILE)
    excel_path = get_allotted_holdings_path()

    if not os.path.exists(excel_path):
        logger.error(f"File not found at '{excel_path}'")
        return

    try:
        wb = Base.safe_load_workbook(excel_path)
    except Exception as e:
        logger.error(f"Error loading workbook '{excel_path}': {e}", exc_info=True)
        return

    file_changed = False

    for user in users:
        uci_user = str(user.get("uci"))
        client_id = user.get("broker_client_id")
        password_user = user.get("password_broker")
        totp_broker = user.get("topt_broker") or user.get("totp_broker")

        if uci_user not in wb.sheetnames:
            continue

        ws = wb[uci_user]

        for r in range(2, ws.max_row + 1):
            security_symbol = ws.cell(r, 1).value
            if not security_symbol or str(security_symbol).strip().lower() in ["none", ""]:
                continue

            security_symbol = str(security_symbol).strip()
            status_cell = ws.cell(r, 8).value if ws.cell(r, 8).value is not None else ws.cell(r, 7).value

            try:
                spl_status = int(status_cell) if status_cell is not None else 0
            except (ValueError, TypeError):
                spl_status = 0

            # Only process rows with special_session_status == NEWLY_ALLOTTED (allotted today / pending order placement)
            if spl_status != SpecialSessionStatus.NEWLY_ALLOTTED:
                continue

            lot_size = ws.cell(r, 2).value or 1
            issue_price = ws.cell(r, 3).value or 0.0
            shares_allocated = ws.cell(r, 4).value or 0
            exchange = str(ws.cell(r, 6).value or "NSE").strip().upper()

            try:
                shares_allocated = int(shares_allocated)
            except (ValueError, TypeError):
                shares_allocated = 0

            try:
                issue_price = float(issue_price)
            except (ValueError, TypeError):
                issue_price = 0.0

            logger.info(f"[User: {uci_user}] Triggering LC Sell Order for '{security_symbol}' | Shares: {shares_allocated} | Exchange: {exchange} | Issue Price: {issue_price}")

            if client_id and password_user and totp_broker:
                try:
                    success, msg = zerodha_sell_lc(
                        user_id=client_id,
                        password=password_user,
                        totp_secret=totp_broker,
                        security_symbol=security_symbol,
                        shares_quantity=shares_allocated,
                        exchange=exchange,
                        issue_price=issue_price,
                    )
                    logger.info(f"Execution Output for {security_symbol}: {msg}")
                    if success:
                        logger.audit(
                            f"[AUDIT] Pre-Open LC Sell Placed: Symbol={security_symbol}, Qty={shares_allocated}, Exchange={exchange}, "
                            f"Issue Price=₹{issue_price}, User={uci_user}"
                        )
                        # Update status to LC_ORDER_PLACED (order placed for special session)
                        status_col = 8 if ws.cell(r, 8).value is not None else 7
                        ws.cell(r, status_col, SpecialSessionStatus.LC_ORDER_PLACED)
                        file_changed = True
                    else:
                        logger.warning(f"LC sell order failed for {security_symbol} ({uci_user}): {msg}")
                except Exception as e:
                    logger.error(f"Error executing LC sell order for {security_symbol} ({uci_user}): {e}", exc_info=True)
            else:
                logger.warning(f"Missing Zerodha credentials for user {uci_user}, cannot place sell order.")

    if file_changed:
        try:
            Base.safe_save_workbook(wb, excel_path)
            logger.info("Successfully updated special_session_status in allotted_holdings.xlsx")
            # Dual-Tier SQLite synchronization
            try:
                import database
                database.import_allotted_holdings_from_excel()
            except Exception as dberr:
                logger.error(f"Error syncing allotted holdings to SQLite: {dberr}", exc_info=True)
        except Exception as e:
            logger.error(f"Error saving workbook '{excel_path}': {e}", exc_info=True)

    wb.close()


if __name__ == "__main__":
    place_lc_sell_orders_for_allotted_today()
